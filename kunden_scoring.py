"""
Kunden-Scoring — Kaufstärke · Zahlungsmoral · Zahlungsgeschwindigkeit
=====================================================================

Bewertet alle Kunden anhand ihrer Ausgangsrechnungen aus Odoo
(account.move, out_invoice / out_refund, gebucht).

Datenquellen (eine von beiden):
  A) Odoo-API (empfohlen) — Umgebungsvariablen:
       ODOO_URL      z. B. https://main.elvinci.opa.as
       ODOO_DB       Datenbankname
       ODOO_USER     Login (E-Mail) des API-Users
       ODOO_API_KEY  API-Key (Odoo: Einstellungen › Konto-Sicherheit › API-Schlüssel), nur Lesezugriff nötig
     Aufruf:  python kunden_scoring.py --odoo
  B) CSV-Export — Pflichtspalten (Semikolon, UTF-8):
       kunde; rechnung; typ (out_invoice|out_refund); rechnungsdatum; faelligkeit;
       netto; brutto; offen; zahlungsdatum
     Aufruf:  python kunden_scoring.py --csv rechnungen.csv

Zahlungsdatum (API): betragsgewichtetes Datum aller Ausgleichsbuchungen
(account.partial.reconcile) der Forderungszeile — Gutschrift-Verrechnungen
zählen NICHT als Zahlung.

Ausgabe: output/YYYY-MM-DD_Kundenbewertung_v1.xlsx (gitignored — enthält Kundendaten)
"""
import argparse
import os
import sys
import io
from datetime import date, timedelta
from pathlib import Path

import pandas as pd

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)

# === Parameter (werden ins Excel-Blatt "Parameter" geschrieben und dort per Formel genutzt) ===
GEWICHT_KAUF = 0.40
GEWICHT_MORAL = 0.35
GEWICHT_SPEED = 0.25
TOLERANZ_TAGE = 3        # Zahlung bis Fälligkeit + 3 Tage gilt als pünktlich
SPEED_NULL_BEI = 60      # Ø Zahltage, ab denen der Speed-Score 0 ist (0 Tage = 100)
ABZUG_MAX = 50           # max. Punktabzug Moral für offene überfällige Posten
MIN_BEZAHLT = 3          # darunter: Datenbasis "gering", kein Gesamtscore
KLASSEN = [(75, 'A'), (55, 'B'), (35, 'C'), (0, 'D')]

PAGE = 2000


# ---------------------------------------------------------------- Odoo-API
class Odoo:
    def __init__(self, url, db, user, key):
        import requests
        self.s = requests.Session()
        self.url = url.rstrip('/') + '/jsonrpc'
        self.db, self.key = db, key
        self.uid = self._call('common', 'authenticate', [db, user, key, {}])
        if not self.uid:
            raise SystemExit('Odoo-Login fehlgeschlagen (ODOO_DB/ODOO_USER/ODOO_API_KEY prüfen).')

    def _call(self, service, method, args):
        r = self.s.post(self.url, json={'jsonrpc': '2.0', 'method': 'call',
                                        'params': {'service': service, 'method': method, 'args': args}},
                        timeout=120)
        r.raise_for_status()
        j = r.json()
        if 'error' in j:
            raise RuntimeError(j['error'].get('data', {}).get('message') or j['error'])
        return j['result']

    def kw(self, model, method, args, **kwargs):
        return self._call('object', 'execute_kw', [self.db, self.uid, self.key, model, method, args, kwargs])

    def search_read_all(self, model, domain, fields):
        out, offset = [], 0
        while True:
            batch = self.kw(model, 'search_read', [domain], fields=fields, offset=offset, limit=PAGE, order='id')
            out += batch
            if len(batch) < PAGE:
                return out
            offset += PAGE

    def read_chunked(self, model, ids, fields):
        ids, out = list(ids), []
        for i in range(0, len(ids), PAGE):
            out += self.kw(model, 'read', [ids[i:i + PAGE]], fields=fields)
        return out


def lade_odoo(start):
    env = {k: os.environ.get(k) for k in ('ODOO_URL', 'ODOO_DB', 'ODOO_USER', 'ODOO_API_KEY')}
    fehlend = [k for k, v in env.items() if not v]
    if fehlend:
        raise SystemExit(f'Fehlende Umgebungsvariablen: {", ".join(fehlend)}')
    o = Odoo(env['ODOO_URL'], env['ODOO_DB'], env['ODOO_USER'], env['ODOO_API_KEY'])

    print(f'[1/4] Rechnungen ab {start} laden …')
    moves = o.search_read_all('account.move', [
        ('move_type', 'in', ['out_invoice', 'out_refund']),
        ('state', '=', 'posted'),
        ('invoice_date', '>=', start.isoformat()),
    ], ['name', 'move_type', 'commercial_partner_id', 'invoice_date', 'invoice_date_due',
        'amount_untaxed_signed', 'amount_total_signed', 'amount_residual_signed', 'payment_state'])
    print(f'      {len(moves):,} Belege')
    if not moves:
        return pd.DataFrame()

    print('[2/4] Forderungszeilen laden …')
    inv_ids = [m['id'] for m in moves if m['move_type'] == 'out_invoice']
    rec_lines = []
    for i in range(0, len(inv_ids), PAGE):
        rec_lines += o.search_read_all('account.move.line', [
            ('move_id', 'in', inv_ids[i:i + PAGE]),
            ('account_type', '=', 'asset_receivable'),
        ], ['move_id', 'matched_credit_ids'])
    line_to_move = {l['id']: l['move_id'][0] for l in rec_lines}

    print('[3/4] Ausgleichsbuchungen laden …')
    partial_ids = {p for l in rec_lines for p in l['matched_credit_ids']}
    partials = o.read_chunked('account.partial.reconcile', partial_ids,
                              ['debit_move_id', 'credit_move_id', 'amount', 'max_date'])
    # Gegenbuchung klassifizieren: Gutschrift (out_refund) ≠ Zahlung
    credit_line_ids = {p['credit_move_id'][0] for p in partials}
    credit_lines = o.read_chunked('account.move.line', credit_line_ids, ['move_id'])
    credit_move_ids = {l['move_id'][0] for l in credit_lines}
    credit_types = {m['id']: m['move_type'] for m in o.read_chunked('account.move', credit_move_ids, ['move_type'])}
    cline_type = {l['id']: credit_types.get(l['move_id'][0]) for l in credit_lines}

    print('[4/4] Zahlungsdatum je Rechnung berechnen …')
    pay = {}  # move_id -> [(betrag, datum)]
    for p in partials:
        if cline_type.get(p['credit_move_id'][0]) in ('out_refund', 'out_invoice'):
            continue
        mid = line_to_move.get(p['debit_move_id'][0])
        if mid:
            pay.setdefault(mid, []).append((p['amount'], pd.Timestamp(p['max_date'])))

    rows = []
    for m in moves:
        zd = None
        if m['id'] in pay and m['payment_state'] in ('paid', 'in_payment'):
            tot = sum(a for a, _ in pay[m['id']])
            if tot > 0:
                base = min(d for _, d in pay[m['id']])
                zd = base + pd.Timedelta(days=sum(a * (d - base).days for a, d in pay[m['id']]) / tot)
        rows.append({
            'kunde': m['commercial_partner_id'][1] if m['commercial_partner_id'] else '(ohne Partner)',
            'rechnung': m['name'], 'typ': m['move_type'],
            'rechnungsdatum': m['invoice_date'], 'faelligkeit': m['invoice_date_due'] or m['invoice_date'],
            'netto': m['amount_untaxed_signed'], 'brutto': m['amount_total_signed'],
            'offen': m['amount_residual_signed'], 'zahlungsdatum': zd,
        })
    return pd.DataFrame(rows)


def lade_csv(pfad):
    df = pd.read_csv(pfad, sep=';', encoding='utf-8-sig')
    pflicht = ['kunde', 'rechnung', 'typ', 'rechnungsdatum', 'faelligkeit', 'netto', 'brutto', 'offen', 'zahlungsdatum']
    fehlend = [c for c in pflicht if c not in df.columns]
    if fehlend:
        raise SystemExit(f'CSV: fehlende Spalten {fehlend}. Vorhanden: {list(df.columns)}')
    return df[pflicht]


# ---------------------------------------------------------------- Scoring
def bewerte(df, stichtag):
    df = df.copy()
    for c in ('rechnungsdatum', 'faelligkeit', 'zahlungsdatum'):
        df[c] = pd.to_datetime(df[c], errors='coerce')
    for c in ('netto', 'brutto', 'offen'):
        df[c] = pd.to_numeric(df[c], errors='coerce').fillna(0.0)
    stichtag = pd.Timestamp(stichtag)

    inv = df[(df['typ'] == 'out_invoice') & (df['brutto'] > 0)].copy()
    bez = inv[inv['zahlungsdatum'].notna()].copy()
    bez['zahltage'] = (bez['zahlungsdatum'] - bez['rechnungsdatum']).dt.days.clip(lower=0)
    bez['verzug'] = (bez['zahlungsdatum'] - bez['faelligkeit']).dt.days.clip(lower=0)
    bez['puenktlich'] = bez['verzug'] <= TOLERANZ_TAGE
    bez['w_zt'] = bez['zahltage'] * bez['brutto']
    bez['w_vz'] = bez['verzug'] * bez['brutto']
    bez['b_p'] = bez['brutto'].where(bez['puenktlich'], 0.0)

    ueb = inv[(inv['offen'] > 0.005) & (inv['faelligkeit'] < stichtag)].copy()
    ueb['tage_ueb'] = (stichtag - ueb['faelligkeit']).dt.days

    k = df.groupby('kunde').agg(umsatz_netto=('netto', 'sum'), umsatz_brutto=('brutto', 'sum'),
                                letzte_rechnung=('rechnungsdatum', 'max'))
    k['anz_rechnungen'] = inv.groupby('kunde').size()
    k['avg_rechnung_netto'] = inv.groupby('kunde')['netto'].mean()
    g = bez.groupby('kunde')
    k['anz_bezahlt'] = g.size()
    k['zahltage_avg'] = g['w_zt'].sum() / g['brutto'].sum()
    k['verzug_avg'] = g['w_vz'].sum() / g['brutto'].sum()
    k['puenktlich_quote'] = g['b_p'].sum() / g['brutto'].sum()
    k['offen_ueberfaellig'] = ueb.groupby('kunde')['offen'].sum()
    k['max_tage_ueberfaellig'] = ueb.groupby('kunde')['tage_ueb'].max()
    k[['anz_rechnungen', 'anz_bezahlt', 'offen_ueberfaellig', 'max_tage_ueberfaellig']] = \
        k[['anz_rechnungen', 'anz_bezahlt', 'offen_ueberfaellig', 'max_tage_ueberfaellig']].fillna(0)

    # Kaufstärke: Perzentil-Rang des Netto-Umsatzes (robust gegen Ausreißer)
    pos = k['umsatz_netto'] > 0
    k['score_kauf'] = 0.0
    k.loc[pos, 'score_kauf'] = k.loc[pos, 'umsatz_netto'].rank(pct=True) * 100

    # Zahlungsmoral: Pünktlichkeitsquote (betragsgewichtet) − Abzug offene Überfälligkeit
    abzug = (k['offen_ueberfaellig'] / k['umsatz_brutto'].where(k['umsatz_brutto'] > 0)).fillna(0) * 100
    k['score_moral'] = (k['puenktlich_quote'] * 100 - abzug.clip(upper=ABZUG_MAX)).clip(lower=0)

    # Geschwindigkeit: 0 Tage = 100, SPEED_NULL_BEI Tage = 0, linear
    k['score_speed'] = ((SPEED_NULL_BEI - k['zahltage_avg']) / SPEED_NULL_BEI * 100).clip(0, 100)

    k['datenbasis'] = k['anz_bezahlt'].apply(lambda n: 'ok' if n >= MIN_BEZAHLT else 'gering')
    voll = k['datenbasis'] == 'ok'
    k['score_gesamt'] = (GEWICHT_KAUF * k['score_kauf'] + GEWICHT_MORAL * k['score_moral']
                         + GEWICHT_SPEED * k['score_speed']).where(voll)
    k['klasse'] = k['score_gesamt'].apply(
        lambda s: next(c for t, c in KLASSEN if s >= t) if pd.notna(s) else 'n/a')
    return k.reset_index().sort_values(['score_gesamt', 'umsatz_netto'], ascending=[False, False],
                                       na_position='last'), df


# ---------------------------------------------------------------- Excel
def schreibe_excel(k, roh, pfad, start, stichtag, quelle):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Bewertung'
    hdr_fill = PatternFill('solid', fgColor='1F2937')
    hdr_font = Font(bold=True, color='FFFFFF')

    cols = [('Rang', None), ('Kunde', 'kunde'), ('Klasse', None), ('Score gesamt', None),
            ('Score Kaufstärke', 'score_kauf'), ('Score Zahlungsmoral', 'score_moral'),
            ('Score Geschwindigkeit', 'score_speed'), ('Datenbasis', 'datenbasis'),
            ('Umsatz netto €', 'umsatz_netto'), ('Rechnungen', 'anz_rechnungen'),
            ('Ø Rechnung netto €', 'avg_rechnung_netto'), ('davon bezahlt', 'anz_bezahlt'),
            ('Ø Zahltage (gew.)', 'zahltage_avg'), ('Ø Verzug Tage (gew.)', 'verzug_avg'),
            ('Pünktlich-Quote', 'puenktlich_quote'), ('Offen überfällig €', 'offen_ueberfaellig'),
            ('Max. Tage überfällig', 'max_tage_ueberfaellig'), ('Letzte Rechnung', 'letzte_rechnung')]
    for j, (h, _) in enumerate(cols, 1):
        c = ws.cell(1, j, h)
        c.fill, c.font = hdr_fill, hdr_font
        c.alignment = Alignment(wrap_text=True, vertical='center')

    for i, (_, r) in enumerate(k.iterrows(), 2):
        for j, (_, key) in enumerate(cols, 1):
            if key is None:
                continue
            v = r[key]
            if pd.isna(v):
                v = None
            elif isinstance(v, pd.Timestamp):
                v = v.to_pydatetime().date()
            elif hasattr(v, 'item'):
                v = v.item()
            ws.cell(i, j, v)
        ws.cell(i, 1, i - 1)
        # Gesamtscore + Klasse als sichtbare Formeln (Gewichte/Schwellen aus Blatt "Parameter")
        ws.cell(i, 4, f'=IF(H{i}="ok",E{i}*Parameter!$B$2+F{i}*Parameter!$B$3+G{i}*Parameter!$B$4,"")')
        ws.cell(i, 3, f'=IF(D{i}="","n/a",IF(D{i}>=Parameter!$B$9,"A",IF(D{i}>=Parameter!$B$10,"B",'
                      f'IF(D{i}>=Parameter!$B$11,"C","D"))))')

    fmt = {4: '0.0', 5: '0.0', 6: '0.0', 7: '0.0', 9: '#,##0.00', 11: '#,##0.00',
           13: '0.0', 14: '0.0', 15: '0.0%', 16: '#,##0.00', 18: 'DD.MM.YYYY'}
    for col, f in fmt.items():
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = f
    widths = [6, 40, 8, 10, 10, 10, 12, 10, 14, 10, 14, 10, 11, 11, 11, 14, 11, 13]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions

    p = wb.create_sheet('Parameter')
    for r in [('Parameter', 'Wert'), ('Gewicht Kaufstärke', GEWICHT_KAUF), ('Gewicht Zahlungsmoral', GEWICHT_MORAL),
              ('Gewicht Geschwindigkeit', GEWICHT_SPEED), ('Toleranz pünktlich (Tage)', TOLERANZ_TAGE),
              ('Speed-Score 0 bei Ø Zahltagen', SPEED_NULL_BEI), ('Max. Abzug Moral', ABZUG_MAX),
              ('Min. bezahlte Rechnungen', MIN_BEZAHLT), ('Schwelle A', KLASSEN[0][0]),
              ('Schwelle B', KLASSEN[1][0]), ('Schwelle C', KLASSEN[2][0])]:
        p.append(r)
    p['A1'].font = p['B1'].font = Font(bold=True)
    p.column_dimensions['A'].width = 32
    p.append([])
    p.append(['Hinweis', 'Gewichte (B2:B4) und Schwellen (B9:B11) wirken per Formel auf Blatt "Bewertung". '
                         'B5–B8 wirken nur bei Neuberechnung im Skript.'])

    m = wb.create_sheet('Methodik')
    for line in [
        f'Kundenbewertung — erstellt {date.today():%d.%m.%Y}, v1',
        f'Datenquelle: {quelle}',
        f'Bewertungszeitraum: Rechnungsdatum {start:%d.%m.%Y} – {stichtag:%d.%m.%Y} (Stichtag)',
        f'Kunden: {len(k):,} | davon mit ausreichender Datenbasis: {(k["datenbasis"] == "ok").sum():,}',
        '',
        'KAUFSTÄRKE (0–100): Perzentil-Rang des Netto-Umsatzes (Rechnungen − Gutschriften) unter allen Kunden mit Umsatz > 0.',
        'ZAHLUNGSMORAL (0–100): Anteil des Brutto-Rechnungsvolumens, der bis Fälligkeit + Toleranz bezahlt wurde, × 100,',
        '   abzüglich (offener überfälliger Betrag / Brutto-Umsatz im Zeitraum × 100), Abzug gedeckelt.',
        'GESCHWINDIGKEIT (0–100): betragsgewichtete Ø Tage Rechnungsdatum → Zahlung; 0 Tage = 100, linear bis 0.',
        'GESAMT: gewichtete Summe (Gewichte s. Parameter). Nur bei ausreichender Datenbasis (≥ Min. bezahlte Rechnungen).',
        '',
        'ZAHLUNGSDATUM: betragsgewichtetes Datum aller Zahlungs-Ausgleiche der Forderungszeile (account.partial.reconcile).',
        '   Verrechnungen mit Gutschriften gelten nicht als Zahlung. Teilzahlungen offener Rechnungen fließen nicht ein.',
        '',
        'NICHT ENTHALTEN / GRENZEN:',
        ' - Rechnungen vor Zeitraumbeginn (auch wenn noch offen) sind nicht erfasst.',
        ' - Mahnstufen, Skonto-Abzüge und Zahlungsbedingungen (z. B. Vorkasse) werden nicht separat bewertet;',
        '   Vorkasse-Kunden erreichen dadurch automatisch hohe Speed-Werte.',
        ' - Kaufstärke ist relativ (Rang), keine absolute Bonitätsaussage.',
        ' - Gewichte und Schwellen sind Startwerte [ANNAHME] und vor Verwendung fachlich freizugeben.',
        '',
        'VERTRAULICH — enthält Kundendaten. Nicht extern weitergeben.',
    ]:
        m.append([line])
    m.column_dimensions['A'].width = 140

    r = wb.create_sheet('Rechnungen')
    roh = roh.copy()
    for c in ('rechnungsdatum', 'faelligkeit', 'zahlungsdatum'):
        roh[c] = pd.to_datetime(roh[c], errors='coerce').dt.date
    r.append(list(roh.columns))
    for row in roh.itertuples(index=False):
        r.append([None if (not isinstance(v, str) and pd.isna(v)) else v for v in row])

    pfad.parent.mkdir(parents=True, exist_ok=True)
    wb.save(pfad)


def main():
    ap = argparse.ArgumentParser(description='Kunden-Scoring aus Odoo-Rechnungen')
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument('--odoo', action='store_true', help='Daten live per Odoo-API laden')
    src.add_argument('--csv', type=Path, help='Rechnungs-CSV (Spalten s. Docstring)')
    ap.add_argument('--monate', type=int, default=12, help='Bewertungszeitraum in Monaten (Standard 12)')
    ap.add_argument('--stichtag', type=date.fromisoformat, default=date.today())
    ap.add_argument('--out', type=Path, default=None)
    a = ap.parse_args()

    start = a.stichtag - timedelta(days=round(a.monate * 30.44))
    if a.odoo:
        roh, quelle = lade_odoo(start), f'Odoo-API {os.environ.get("ODOO_URL")} (account.move, gebucht)'
    else:
        roh, quelle = lade_csv(a.csv), f'CSV {a.csv.name}'
        rd = pd.to_datetime(roh['rechnungsdatum'], errors='coerce')
        roh = roh[(rd >= pd.Timestamp(start)) & (rd <= pd.Timestamp(a.stichtag))]
    if roh.empty:
        raise SystemExit('Keine Rechnungen im Zeitraum gefunden.')

    k, roh = bewerte(roh, a.stichtag)
    out = a.out or Path('output') / f'{date.today():%Y-%m-%d}_Kundenbewertung_v1.xlsx'
    schreibe_excel(k, roh, out, start, a.stichtag, quelle)

    print(f'\nKunden bewertet: {len(k):,}  (Datenbasis ok: {(k["datenbasis"] == "ok").sum():,})')
    print('Klassenverteilung:', k['klasse'].value_counts().to_dict())
    print(f'Datei: {out}')


if __name__ == '__main__':
    main()
